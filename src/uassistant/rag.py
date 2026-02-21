# src/uassistant/rag.py
"""Module RAG local : indexation et recherche sémantique sans dépendances cloud."""
import os
import glob
import numpy as np
from PyPDF2 import PdfReader
import openpyxl
from docx import Document
from sentence_transformers import SentenceTransformer
from .utils import resource_path

class LocalRAG:
    def __init__(self, model_path: str = None):
        model_path = model_path or resource_path("embedding_model")
        if not os.path.exists(model_path):
            raise FileNotFoundError(f"Modèle non trouvé : {model_path}")
        self.embedding_model = SentenceTransformer(
            model_path,
            local_files_only=True,
            trust_remote_code=False
        )
        self.chunks: list[str] = []
        self.embeddings: np.ndarray = None
        self.sources: list[str] = []
    
    def extract_text(self, file_path: str) -> str:
        ext = os.path.splitext(file_path)[1].lower()
        text = ""
        try:
            if ext == ".pdf":
                reader = PdfReader(file_path)
                for page in reader.pages:
                    content = page.extract_text()
                    if content: text += content + "\n"
            elif ext == ".xlsx":
                wb = openpyxl.load_workbook(file_path, read_only=True)
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        text += " ".join(str(c) for c in row if c) + "\n"
            elif ext == ".txt":
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    text = f.read()
            elif ext == ".docx":
                doc = Document(file_path)
                for para in doc.paragraphs:
                    text += para.text + "\n"
        except Exception as e:
            print(f"FEHLER beim Lesen von {file_path}: {e}")
        return text
    
    def split_text(self, text: str, chunk_size: int = 500, overlap: int = 50) -> list[str]:
        if not text.strip(): return []
        # Respecte les séparateurs === si présents
        if "===" in text:
            sections = [s.strip() for s in text.split("===") if s.strip()]
            chunks = []
            for section in sections:
                chunks.extend(self._split_section(section, chunk_size, overlap))
            return chunks
        return self._split_section(text, chunk_size, overlap)
    
    def _split_section(self, text: str, chunk_size: int, overlap: int) -> list[str]:
        if len(text) <= chunk_size: return [text]
        chunks, start = [], 0
        while start < len(text):
            end = min(start + chunk_size, len(text))
            if end < len(text):
                while end > start and text[end] not in " \n.!?": end -= 1
            chunk = text[start:end].strip()
            if chunk: chunks.append(chunk)
            start = end - overlap if end - overlap > start else end
        return chunks
    
    def ingest_documents(self, data_dir: str = "./data"):
        os.makedirs(data_dir, exist_ok=True)
        files = glob.glob(os.path.join(data_dir, "**", "*.*"), recursive=True)
        files = [f for f in files if f.endswith((".pdf", ".xlsx", ".txt", ".docx"))]
        
        all_chunks, all_sources = [], []
        for file in files:
            text = self.extract_text(file)
            if text.strip():
                chunks = self.split_text(text)
                all_chunks.extend(chunks)
                all_sources.extend([os.path.basename(file)] * len(chunks))
        
        if all_chunks:
            embeddings = self.embedding_model.encode(all_chunks, show_progress_bar=False)
            self.chunks = all_chunks
            self.embeddings = np.array(embeddings)
            self.sources = all_sources
            print(f"✅ {len(all_chunks)} chunks indexés depuis {len(files)} fichiers.")
        else:
            print("⚠️  Aucun contenu lisible trouvé dans ./data/")
    
    def retrieve(self, query: str, k: int = 5, min_score: float = 0.25) -> str:
        if not self.chunks or self.embeddings is None: return ""
        query_emb = self.embedding_model.encode([query])[0]
        scores = np.dot(self.embeddings, query_emb) / (
            np.linalg.norm(self.embeddings, axis=1) * np.linalg.norm(query_emb) + 1e-8
        )
        top_idx = np.argsort(scores)[::-1]
        results = []
        for idx in top_idx:
            if scores[idx] >= min_score and len(results) < k:
                results.append(f"[{self.sources[idx]}] {self.chunks[idx]}")
            elif len(results) >= k: break
        return "\n".join(results) if results else ""